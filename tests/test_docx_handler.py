from pathlib import Path

from docx import Document

from app.documents import DocxHandler


def test_replaces_markers_in_body_and_table(tmp_path: Path):
    template = tmp_path / "template.docx"
    output = tmp_path / "output.docx"

    document = Document()
    document.add_paragraph("ФИО: {{ФИО}}")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Адрес: {{Адрес}}"
    document.save(template)

    DocxHandler().render(
        template,
        output,
        {"{{ФИО}}": "Аксенов Иван Петрович", "{{Адрес}}": "Москва"},
    )

    result = Document(output)
    assert "Аксенов Иван Петрович" in result.paragraphs[0].text
    assert "Москва" in result.tables[0].cell(0, 0).text


def test_replaces_all_occurrences_and_preserves_unmatched_marker(tmp_path: Path):
    template = tmp_path / "template.docx"
    output = tmp_path / "output.docx"

    document = Document()
    paragraph = document.add_paragraph("{{ФИО}} / {{ФИО}} / {{Неизвестно}}")
    paragraph.runs[0].bold = True
    document.save(template)

    DocxHandler().render(
        template,
        output,
        {"{{ФИО}}": "Аксенов Иван Петрович"},
    )

    result = Document(output)
    assert result.paragraphs[0].text == (
        "Аксенов Иван Петрович / Аксенов Иван Петрович / {{Неизвестно}}"
    )
    assert result.paragraphs[0].runs[0].bold is True


def test_empty_value_leaves_marker_unchanged(tmp_path: Path):
    template = tmp_path / "template.docx"
    output = tmp_path / "output.docx"

    document = Document()
    document.add_paragraph("Телефон: {{Телефон}}")
    document.save(template)

    DocxHandler().render(template, output, {"{{Телефон}}": ""})

    result = Document(output)
    assert result.paragraphs[0].text == "Телефон: {{Телефон}}"


def test_replaces_markers_in_headers_and_footers(tmp_path: Path):
    template = tmp_path / "template.docx"
    output = tmp_path / "output.docx"

    document = Document()
    section = document.sections[0]
    section.header.paragraphs[0].text = "Группа: {{Группа}}"
    section.footer.paragraphs[0].text = "ФИО: {{ФИО}}"
    document.add_paragraph("Body")
    document.save(template)

    DocxHandler().render(
        template,
        output,
        {"{{Группа}}": "211_26", "{{ФИО}}": "Аксенов Иван Петрович"},
    )

    result = Document(output)
    assert result.sections[0].header.paragraphs[0].text == "Группа: 211_26"
    assert result.sections[0].footer.paragraphs[0].text == "ФИО: Аксенов Иван Петрович"
