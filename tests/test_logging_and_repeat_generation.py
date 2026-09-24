from pathlib import Path

from docx import Document

from app.generation import Generator
from app.paths import ensure_runtime_directories


def _write_registry(path: Path, rows: list[list[str]]) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    workbook.save(path)


def _write_template(path: Path, text: str) -> None:
    document = Document()
    document.add_paragraph(text)
    document.save(path)


def test_generation_writes_dedicated_log(tmp_path: Path):
    app_root = tmp_path / "app"
    registry = tmp_path / "registry.xlsx"
    _write_registry(
        registry,
        [["{{ФИО}}"], ["Аксенов Иван Петрович"]],
    )

    templates, logs, _results = ensure_runtime_directories(app_root)
    _write_template(templates / "Справка.docx", "ФИО: {{ФИО}}")

    stats = Generator().generate(registry, "101", app_root)

    assert stats.generated == 1
    log_files = list(logs.glob("Генерация_*.log"))
    assert len(log_files) == 1
    log_text = log_files[0].read_text(encoding="utf-8")
    assert "Excel:" in log_text
    assert "Группа: 101" in log_text
    assert "Создано:" in log_text
    assert "Итоги:" in log_text


def test_empty_registry_still_writes_log_but_no_results(tmp_path: Path):
    app_root = tmp_path / "app"
    registry = tmp_path / "registry.xlsx"
    _write_registry(registry, [["{{ФИО}}"], [None]])

    _templates, logs, results = ensure_runtime_directories(app_root)
    stats = Generator().generate(registry, "101", app_root)

    assert stats.students == 0
    assert list(logs.glob("Генерация_*.log"))
    assert not (results / "101").exists()


def test_second_run_overwrites_normal_output_and_preserves_error_copy(
    tmp_path: Path,
):
    app_root = tmp_path / "app"
    registry = tmp_path / "registry.xlsx"
    _write_registry(registry, [["{{ФИО}}"], ["Аксенов Иван Петрович"]])

    templates, _logs, results = ensure_runtime_directories(app_root)
    template = templates / "Справка.docx"
    _write_template(template, "Версия 1: {{ФИО}}")

    Generator().generate(registry, "101", app_root)
    output = results / "101" / "Аксенов Иван Петрович" / "Аксенов_Справка.docx"
    assert "Версия 1" in Document(output).paragraphs[0].text

    _write_template(template, "Версия 2: {{ФИО}}")
    Generator().generate(registry, "101", app_root)
    assert "Версия 2" in Document(output).paragraphs[0].text

    error_copy = (
        results
        / "101"
        / "Аксенов Иван Петрович"
        / "Аксенов_Справка_ОШИБКА.docx"
    )
    error_copy.write_bytes(b"manual diagnostic file")

    # A later successful run must not delete a pre-existing _ОШИБКА artifact.
    Generator().generate(registry, "101", app_root)
    assert error_copy.read_bytes() == b"manual diagnostic file"
