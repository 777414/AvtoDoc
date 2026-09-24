from pathlib import Path

from docx import Document

from app.generation import Generator
from app.paths import ensure_runtime_directories


def _write_registry(path: Path, rows: list[list[str]]) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    workbook.save(path)


def _write_template(path: Path, text: str) -> None:
    document = Document()
    document.add_paragraph(text)
    document.save(path)


def test_generates_every_student_template_and_creates_runtime_directories(
    tmp_path: Path,
):
    app_root = tmp_path / "app"
    registry = tmp_path / "registry.xlsx"
    _write_registry(
        registry,
        [
            ["{{ФИО}}", "{{Группа}}"],
            ["Аксенов Иван Петрович", "101"],
            ["Петров Сергей Сергеевич", "101"],
        ],
    )

    templates, logs, results = ensure_runtime_directories(app_root)
    _write_template(templates / "Справка.docx", "ФИО: {{ФИО}}")
    _write_template(templates / "Договор.docx", "Группа: {{Группа}}")

    stats = Generator().generate(registry, "101", app_root)

    assert stats.students == 2
    assert stats.templates == 2
    assert stats.generated == 4
    assert stats.errors == 0
    assert logs.is_dir()
    assert results.is_dir()

    first = results / "101" / "Аксенов Иван Петрович"
    second = results / "101" / "Петров Сергей Сергеевич"
    assert (first / "Аксенов_Справка.docx").is_file()
    assert (first / "Аксенов_Договор.docx").is_file()
    assert (second / "Петров_Справка.docx").is_file()
    assert (second / "Петров_Договор.docx").is_file()

    assert "Аксенов Иван Петрович" in Document(
        first / "Аксенов_Справка.docx"
    ).paragraphs[0].text


def test_no_students_creates_no_result_directory(tmp_path: Path):
    app_root = tmp_path / "app"
    registry = tmp_path / "registry.xlsx"
    _write_registry(registry, [["{{ФИО}}"], [None]])

    stats = Generator().generate(registry, "101", app_root)

    assert stats.students == 0
    assert stats.generated == 0
    assert not (app_root / "Результат").exists()


def test_unsupported_templates_are_ignored_but_student_folders_are_created(
    tmp_path: Path,
):
    app_root = tmp_path / "app"
    registry = tmp_path / "registry.xlsx"
    _write_registry(registry, [["{{ФИО}}"], ["Аксенов Иван Петрович"]])

    templates, _logs, results = ensure_runtime_directories(app_root)
    (templates / "readme.txt").write_text("manual", encoding="utf-8")

    stats = Generator().generate(registry, "101", app_root)

    assert stats.templates == 0
    assert stats.generated == 0
    assert stats.skipped_templates == ["readme.txt"]
    assert (results / "101" / "Аксенов Иван Петрович").is_dir()


def test_failed_template_creates_error_copy_and_continues(tmp_path: Path):
    app_root = tmp_path / "app"
    registry = tmp_path / "registry.xlsx"
    _write_registry(
        registry,
        [
            ["{{ФИО}}"],
            ["Аксенов Иван Петрович"],
        ],
    )

    templates, _logs, results = ensure_runtime_directories(app_root)
    broken = templates / "Плохой.docx"
    broken.write_text("not a docx", encoding="utf-8")

    stats = Generator().generate(registry, "101", app_root)

    assert stats.errors == 1
    assert stats.generated == 0
    error_copy = (
        results
        / "101"
        / "Аксенов Иван Петрович"
        / "Аксенов_Плохой_ОШИБКА.docx"
    )
    assert error_copy.read_bytes() == broken.read_bytes()
    assert broken.read_text(encoding="utf-8") == "not a docx"
