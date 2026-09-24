from pathlib import Path

from docx import Document

from app.generation import Generator
from app.paths import ensure_runtime_directories


def test_generator_reports_progress(tmp_path: Path):
    app_root = tmp_path / "app"
    registry = tmp_path / "registry.xlsx"

    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["{{ФИО}}"])
    sheet.append(["Аксенов Иван Петрович"])
    workbook.save(registry)

    templates, _logs, _results = ensure_runtime_directories(app_root)
    document = Document()
    document.add_paragraph("{{ФИО}}")
    document.save(templates / "Справка.docx")

    messages: list[str] = []
    stats = Generator().generate(registry, "101", app_root, progress=messages.append)

    assert stats.generated == 1
    assert any(message.startswith("Создано:") for message in messages)
    assert messages[-1] == "=== Генерация завершена ==="
