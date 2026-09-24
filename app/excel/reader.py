"""Excel registry reader for AvtoDoc."""

from pathlib import Path

from .models import StudentRecord


SUPPORTED_SUFFIXES = {".xlsx", ".xls"}


class ExcelReadError(Exception):
    """Raised when an Excel registry cannot be read."""


def read_registry(path: str | Path) -> tuple[list[str], list[StudentRecord]]:
    """Read an Excel registry.

    The first row is treated as marker headers. Every subsequent row is a
    student unless all its cells are empty. Cell values are converted to text
    without application-level date/number formatting.
    """
    file_path = Path(path)

    if file_path.suffix.lower() not in SUPPORTED_SUFFIXES:
        raise ExcelReadError(f"Неподдерживаемый формат Excel: {file_path.suffix}")

    try:
        if file_path.suffix.lower() == ".xlsx":
            return _read_xlsx(file_path)

        return _read_xls(file_path)
    except ExcelReadError:
        raise
    except Exception as exc:
        raise ExcelReadError(str(exc)) from exc


def _read_xlsx(path: Path) -> tuple[list[str], list[StudentRecord]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        headers = _read_headers(rows)
        students = _read_students(headers, rows)
        return headers, students
    finally:
        workbook.close()


def _read_xls(path: Path) -> tuple[list[str], list[StudentRecord]]:
    import xlrd

    workbook = xlrd.open_workbook(path, on_demand=True)
    try:
        worksheet = workbook.sheet_by_index(0)
        if worksheet.nrows == 0:
            return [], []

        headers = [_cell_to_text(worksheet.cell_value(0, col)) for col in range(worksheet.ncols)]
        students: list[StudentRecord] = []

        for row_index in range(1, worksheet.nrows):
            values = [
                _cell_to_text(worksheet.cell_value(row_index, col))
                if col < worksheet.ncols
                else ""
                for col in range(len(headers))
            ]
            if any(value != "" for value in values):
                students.append(StudentRecord(dict(zip(headers, values))))

        return headers, students
    finally:
        workbook.release_resources()


def _read_headers(rows) -> list[str]:
    try:
        first_row = next(rows)
    except StopIteration:
        return []

    return [_cell_to_text(value) for value in first_row]


def _read_students(headers: list[str], rows) -> list[StudentRecord]:
    students: list[StudentRecord] = []

    for row in rows:
        values = [_cell_to_text(value) for value in row]
        values = values[: len(headers)] + [""] * max(0, len(headers) - len(values))

        if any(value != "" for value in values):
            students.append(StudentRecord(dict(zip(headers, values))))

    return students


def _cell_to_text(value) -> str:
    if value is None:
        return ""

    return str(value)
