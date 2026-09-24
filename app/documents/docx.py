"""DOCX template handler for AvtoDoc."""

from pathlib import Path
from typing import Mapping

from docx import Document

from .base import DocumentHandler, DocumentProcessingError


class DocxHandler(DocumentHandler):
    supported_suffixes = frozenset({".docx"})

    def render(
        self,
        template_path: Path,
        output_path: Path,
        values: Mapping[str, str],
    ) -> None:
        try:
            document = Document(template_path)
            self._replace_document(document, values)
            document.save(output_path)
        except Exception as exc:
            raise DocumentProcessingError(str(exc)) from exc

    def _replace_document(self, document: Document, values: Mapping[str, str]) -> None:
        for paragraph in document.paragraphs:
            self._replace_paragraph(paragraph, values)

        for table in document.tables:
            self._replace_table(table, values)

        for section in document.sections:
            self._replace_paragraphs(section.header.paragraphs, values)
            self._replace_paragraphs(section.footer.paragraphs, values)

    def _replace_table(self, table, values: Mapping[str, str]) -> None:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    self._replace_paragraph(paragraph, values)
                for nested_table in cell.tables:
                    self._replace_table(nested_table, values)

    def _replace_paragraphs(self, paragraphs, values: Mapping[str, str]) -> None:
        for paragraph in paragraphs:
            self._replace_paragraph(paragraph, values)

    @staticmethod
    def _replace_paragraph(paragraph, values: Mapping[str, str]) -> None:
        for run in paragraph.runs:
            original = run.text
            if not original:
                continue

            updated = original
            for marker, value in values.items():
                if value != "":
                    updated = updated.replace(marker, value)

            run.text = updated
